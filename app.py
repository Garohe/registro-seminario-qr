import os
import csv
import io
import sqlite3
import uuid
from datetime import datetime

import qrcode
from flask import (
    Flask, render_template, request, redirect, url_for,
    flash, jsonify, send_file, send_from_directory
)
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm, mm
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from PIL import Image

BASE_DIR = os.path.abspath(os.path.dirname(__file__))

app = Flask(__name__)
app.secret_key = 'seminario2026-registro-qr-secretkey'

DB_PATH = os.path.join(BASE_DIR, 'database.db')
CREDENCIALES_DIR = os.path.join(BASE_DIR, 'credenciales')
os.makedirs(CREDENCIALES_DIR, exist_ok=True)

QR_TEMP_DIR = os.path.join(BASE_DIR, 'qr_temp')
os.makedirs(QR_TEMP_DIR, exist_ok=True)


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    conn.execute('''
        CREATE TABLE IF NOT EXISTS invitados (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL,
            codigo TEXT UNIQUE NOT NULL,
            asistencia INTEGER DEFAULT 0,
            hora_registro TEXT,
            lote TEXT DEFAULT ''
        )
    ''')
    conn.commit()
    conn.close()


init_db()


@app.route('/')
def index():
    conn = get_db()
    total = conn.execute('SELECT COUNT(*) FROM invitados').fetchone()[0]
    presentes = conn.execute('SELECT COUNT(*) FROM invitados WHERE asistencia = 1').fetchone()[0]
    conn.close()
    return render_template('index.html', total=total, presentes=presentes)


@app.route('/importar', methods=['GET', 'POST'])
def importar():
    if request.method == 'POST':
        archivo = request.files.get('archivo')
        if not archivo or archivo.filename == '':
            flash('No se seleccionó ningún archivo', 'error')
            return redirect(url_for('importar'))

        limpiar = request.form.get('limpiar') == '1'
        conn = get_db()

        if limpiar:
            conn.execute('DELETE FROM invitados')
            conn.commit()

        max_code = conn.execute('SELECT MAX(CAST(SUBSTR(codigo, 5) AS INTEGER)) FROM invitados').fetchone()[0]
        counter = (max_code or 0) + 1

        nombres = []
        filename = archivo.filename.lower()

        if filename.endswith('.csv'):
            stream = io.TextIOWrapper(archivo.stream, encoding='utf-8-sig')
            reader = csv.reader(stream)
            header = next(reader, None)
            if header:
                nombre_idx = 0
                for i, col in enumerate(header):
                    if col.strip().lower() == 'nombre':
                        nombre_idx = i
                        break
                for row in reader:
                    if row and row[nombre_idx].strip():
                        nombres.append(row[nombre_idx].strip())

        elif filename.endswith('.xlsx') or filename.endswith('.xls'):
            wb = load_workbook(archivo, read_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if rows:
                header = rows[0]
                nombre_idx = 0
                for i, col in enumerate(header):
                    if col and str(col).strip().lower() == 'nombre':
                        nombre_idx = i
                        break
                for row in rows[1:]:
                    if row[nombre_idx] and str(row[nombre_idx]).strip():
                        nombres.append(str(row[nombre_idx]).strip())
        else:
            flash('Formato no soportado. Use CSV o Excel (.xlsx)', 'error')
            return redirect(url_for('importar'))

        insertados = 0
        for nombre in nombres:
            codigo = f'INV-{counter:04d}'
            try:
                conn.execute(
                    'INSERT INTO invitados (nombre, codigo) VALUES (?, ?)',
                    (nombre, codigo)
                )
                counter += 1
                insertados += 1
            except sqlite3.IntegrityError:
                pass

        conn.commit()
        conn.close()
        flash(f'Se importaron {insertados} invitados exitosamente', 'success')
        return redirect(url_for('invitados'))

    return render_template('importar.html')


@app.route('/invitados')
def invitados():
    buscar = request.args.get('buscar', '')
    lote_filtro = request.args.get('lote', '')
    conn = get_db()

    query = 'SELECT * FROM invitados WHERE 1=1'
    params = []

    if buscar:
        query += ' AND nombre LIKE ?'
        params.append(f'%{buscar}%')
    if lote_filtro:
        query += ' AND lote = ?'
        params.append(lote_filtro)

    query += ' ORDER BY id'
    invitados_list = conn.execute(query, params).fetchall()
    lotes = conn.execute('SELECT DISTINCT lote FROM invitados WHERE lote != "" ORDER BY lote').fetchall()
    conn.close()
    return render_template('invitados.html', invitados=invitados_list, buscar=buscar,
                           lote_filtro=lote_filtro, lotes=lotes)


@app.route('/invitados/editar/<int:id>', methods=['POST'])
def editar_invitado(id):
    nombre = request.form.get('nombre', '').strip()
    lote = request.form.get('lote', '').strip()
    conn = get_db()
    conn.execute('UPDATE invitados SET nombre = ?, lote = ? WHERE id = ?', (nombre, lote, id))
    conn.commit()
    conn.close()
    flash('Invitado actualizado', 'success')
    return redirect(url_for('invitados'))


@app.route('/invitados/eliminar/<int:id>', methods=['POST'])
def eliminar_invitado(id):
    conn = get_db()
    conn.execute('DELETE FROM invitados WHERE id = ?', (id,))
    conn.commit()
    conn.close()
    flash('Invitado eliminado', 'success')
    return redirect(url_for('invitados'))


@app.route('/borrar-todo', methods=['POST'])
def borrar_todo():
    conn = get_db()
    conn.execute('DELETE FROM invitados')
    conn.execute('DELETE FROM sqlite_sequence WHERE name = "invitados"')
    conn.commit()
    conn.close()
    flash('Se borraron todos los invitados y asistencias', 'success')
    return redirect(url_for('importar'))


@app.route('/credenciales')
def credenciales():
    conn = get_db()
    lotes = conn.execute('SELECT DISTINCT lote FROM invitados WHERE lote != "" ORDER BY lote').fetchall()
    total = conn.execute('SELECT COUNT(*) FROM invitados').fetchone()[0]
    conn.close()
    return render_template('credenciales.html', lotes=lotes, total=total)


@app.route('/credenciales/generar')
def generar_credenciales():
    lote = request.args.get('lote', '')
    conn = get_db()

    if lote:
        invitados_list = conn.execute(
            'SELECT * FROM invitados WHERE lote = ? ORDER BY id', (lote,)
        ).fetchall()
    else:
        invitados_list = conn.execute('SELECT * FROM invitados ORDER BY id').fetchall()

    conn.close()

    if not invitados_list:
        flash('No hay invitados para generar credenciales', 'error')
        return redirect(url_for('credenciales'))

    nombre_archivo = f'credenciales_{lote or "todas"}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf'
    ruta_pdf = os.path.join(CREDENCIALES_DIR, nombre_archivo)

    VERDE_MENTA = colors.HexColor('#2ecc71')
    VERDE_OSCURO = colors.HexColor('#27ae60')
    VERDE_CLARO = colors.HexColor('#a8e6cf')
    GRIS_TEXTO = colors.HexColor('#2c3e50')
    GRIS_SUAVE = colors.HexColor('#95a5a6')

    page_w, page_h = letter
    c = canvas.Canvas(ruta_pdf, pagesize=letter)

    badge_w = page_w - 2 * cm
    badge_h = (page_h - 3 * cm) / 2

    for i, inv in enumerate(invitados_list):
        pos_in_page = i % 2
        if pos_in_page == 0 and i > 0:
            c.showPage()

        x = cm
        y = page_h - cm - (pos_in_page * (badge_h + cm))

        c.setStrokeColor(colors.HexColor('#dfe6e9'))
        c.setLineWidth(0.5)
        c.roundRect(x, y - badge_h, badge_w, badge_h, 12)

        header_h = 3.2 * cm
        c.saveState()
        p = c.beginPath()
        p.roundRect(x, y - header_h, badge_w, header_h, 12)
        c.clipPath(p, stroke=0)
        c.setFillColor(VERDE_MENTA)
        c.rect(x, y - header_h, badge_w, header_h, fill=1, stroke=0)
        c.restoreState()
        c.setFillColor(VERDE_MENTA)
        c.rect(x, y - header_h, badge_w, 0.8 * cm, fill=1, stroke=0)

        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 22)
        c.drawCentredString(x + badge_w / 2, y - 1.4 * cm, "SEMINARIO 2026")
        c.setFont("Helvetica", 11)
        c.drawCentredString(x + badge_w / 2, y - 2.1 * cm, "CREDENCIAL DE ACCESO")

        stripe_y = y - header_h - 0.15 * cm
        c.setFillColor(VERDE_OSCURO)
        c.rect(x + 2 * cm, stripe_y, badge_w - 4 * cm, 3, fill=1, stroke=0)

        c.setFillColor(GRIS_TEXTO)
        nombre = inv['nombre']
        font_size = 26
        if len(nombre) > 25:
            font_size = 22
        if len(nombre) > 35:
            font_size = 18
        if len(nombre) > 45:
            font_size = 14
        c.setFont("Helvetica-Bold", font_size)
        nombre_y = y - header_h - 1.5 * cm
        c.drawCentredString(x + badge_w / 2, nombre_y, nombre)

        if inv['lote']:
            c.setFont("Helvetica-Bold", 12)
            c.setFillColor(VERDE_OSCURO)
            lote_y = nombre_y - 0.8 * cm
            lote_text = f"LOTE {inv['lote']}"
            tw = c.stringWidth(lote_text, "Helvetica-Bold", 12)
            pill_x = x + (badge_w - tw) / 2 - 0.4 * cm
            pill_w = tw + 0.8 * cm
            c.setFillColor(VERDE_CLARO)
            c.roundRect(pill_x, lote_y - 0.15 * cm, pill_w, 0.6 * cm, 3, fill=1, stroke=0)
            c.setFillColor(VERDE_OSCURO)
            c.drawCentredString(x + badge_w / 2, lote_y, lote_text)

        qr_path = os.path.join(QR_TEMP_DIR, f"{inv['codigo']}.png")
        qr = qrcode.QRCode(version=1, box_size=10, border=2)
        qr.add_data(inv['codigo'])
        qr.make(fit=True)
        qr_img = qr.make_image(fill_color="#2c3e50", back_color="white")
        qr_img.save(qr_path)

        qr_size = 4.5 * cm
        qr_x = x + (badge_w - qr_size) / 2
        qr_y = y - badge_h + 2.2 * cm
        c.setStrokeColor(VERDE_CLARO)
        c.setLineWidth(2)
        c.roundRect(qr_x - 0.2 * cm, qr_y - 0.2 * cm, qr_size + 0.4 * cm, qr_size + 0.4 * cm, 6)
        c.drawImage(qr_path, qr_x, qr_y, qr_size, qr_size)

        c.setFont("Helvetica", 9)
        c.setFillColor(GRIS_SUAVE)
        c.drawCentredString(x + badge_w / 2, y - badge_h + 1.4 * cm, inv['codigo'])

        footer_y = y - badge_h
        c.saveState()
        p = c.beginPath()
        p.roundRect(x, footer_y, badge_w, 1 * cm, 12)
        c.clipPath(p, stroke=0)
        c.setFillColor(VERDE_MENTA)
        c.rect(x, footer_y, badge_w, 1 * cm, fill=1, stroke=0)
        c.restoreState()
        c.setFillColor(VERDE_MENTA)
        c.rect(x, footer_y + 0.4 * cm, badge_w, 0.6 * cm, fill=1, stroke=0)
        c.setFont("Helvetica", 7)
        c.setFillColor(colors.white)
        c.drawCentredString(x + badge_w / 2, footer_y + 0.3 * cm, "Presente esta credencial en el punto de acceso")

    c.save()

    for f in os.listdir(QR_TEMP_DIR):
        os.remove(os.path.join(QR_TEMP_DIR, f))

    return send_file(ruta_pdf, as_attachment=True, download_name=nombre_archivo)


@app.route('/escanear')
def escanear():
    return render_template('escanear.html')


@app.route('/dashboard')
def dashboard():
    return render_template('dashboard.html')


# ---- API ----

@app.route('/api/registrar', methods=['POST'])
def api_registrar():
    data = request.get_json()
    if not data or 'codigo' not in data:
        return jsonify({'status': 'error', 'message': 'Código no proporcionado'}), 400

    codigo = data['codigo'].strip()
    conn = get_db()
    inv = conn.execute('SELECT * FROM invitados WHERE codigo = ?', (codigo,)).fetchone()

    if not inv:
        conn.close()
        return jsonify({'status': 'error', 'message': 'Código no encontrado'}), 404

    if inv['asistencia'] == 1:
        conn.close()
        return jsonify({
            'status': 'duplicado',
            'message': f'{inv["nombre"]} ya fue registrado',
            'nombre': inv['nombre'],
            'hora': inv['hora_registro']
        })

    hora = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn.execute(
        'UPDATE invitados SET asistencia = 1, hora_registro = ? WHERE codigo = ?',
        (hora, codigo)
    )
    conn.commit()
    conn.close()

    return jsonify({
        'status': 'ok',
        'message': f'{inv["nombre"]} registrado exitosamente',
        'nombre': inv['nombre'],
        'hora': hora
    })


@app.route('/api/estadisticas')
def api_estadisticas():
    conn = get_db()
    total = conn.execute('SELECT COUNT(*) FROM invitados').fetchone()[0]
    presentes = conn.execute('SELECT COUNT(*) FROM invitados WHERE asistencia = 1').fetchone()[0]

    por_lote = conn.execute('''
        SELECT
            CASE WHEN lote = '' THEN 'Sin lote' ELSE lote END as lote,
            COUNT(*) as total,
            SUM(CASE WHEN asistencia = 1 THEN 1 ELSE 0 END) as presentes
        FROM invitados
        GROUP BY lote
        ORDER BY lote
    ''').fetchall()

    ultimos = conn.execute('''
        SELECT nombre, codigo, lote, hora_registro
        FROM invitados
        WHERE asistencia = 1
        ORDER BY hora_registro DESC
        LIMIT 20
    ''').fetchall()

    conn.close()

    return jsonify({
        'total': total,
        'presentes': presentes,
        'ausentes': total - presentes,
        'porcentaje': round((presentes / total * 100), 1) if total > 0 else 0,
        'por_lote': [dict(r) for r in por_lote],
        'ultimos': [dict(r) for r in ultimos]
    })


@app.route('/exportar')
def exportar_excel():
    conn = get_db()
    invitados_list = conn.execute('SELECT * FROM invitados ORDER BY lote, nombre').fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencia Seminario"

    header_fill = PatternFill(start_color="2ecc71", end_color="2ecc71", fill_type="solid")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    verde = PatternFill(start_color="d4edda", end_color="d4edda", fill_type="solid")
    rojo = PatternFill(start_color="f8d7da", end_color="f8d7da", fill_type="solid")
    verde_texto = Font(name="Calibri", color="155724", size=11)
    rojo_texto = Font(name="Calibri", color="721c24", size=11)
    borde = Border(
        left=Side(style="thin", color="dee2e6"),
        right=Side(style="thin", color="dee2e6"),
        top=Side(style="thin", color="dee2e6"),
        bottom=Side(style="thin", color="dee2e6"),
    )
    centrado = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('A1:F1')
    ws['A1'] = "REPORTE DE ASISTENCIA — SEMINARIO 2026"
    ws['A1'].font = Font(name="Calibri", bold=True, size=14, color="2c3e50")
    ws['A1'].alignment = Alignment(horizontal="center")

    total = len(invitados_list)
    presentes = sum(1 for i in invitados_list if i['asistencia'] == 1)
    ws.merge_cells('A2:F2')
    ws['A2'] = f"Total: {total}  |  Presentes: {presentes}  |  Ausentes: {total - presentes}  |  Asistencia: {round(presentes/total*100, 1) if total else 0}%"
    ws['A2'].font = Font(name="Calibri", size=11, color="666666")
    ws['A2'].alignment = Alignment(horizontal="center")

    headers = ["#", "Código", "Nombre", "Lote", "Asistencia", "Hora de Registro"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = centrado
        cell.border = borde

    for idx, inv in enumerate(invitados_list, 1):
        row = idx + 4
        ws.cell(row=row, column=1, value=idx).alignment = centrado
        ws.cell(row=row, column=2, value=inv['codigo']).alignment = centrado
        ws.cell(row=row, column=3, value=inv['nombre'])
        ws.cell(row=row, column=4, value=inv['lote'] or "").alignment = centrado

        asistio = inv['asistencia'] == 1
        cell_asist = ws.cell(row=row, column=5, value="PRESENTE" if asistio else "AUSENTE")
        cell_asist.alignment = centrado
        cell_asist.font = verde_texto if asistio else rojo_texto
        cell_asist.fill = verde if asistio else rojo

        ws.cell(row=row, column=6, value=inv['hora_registro'] or "").alignment = centrado

        for col in range(1, 7):
            ws.cell(row=row, column=col).border = borde

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 35
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 22

    nombre_archivo = f'asistencia_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    ruta = os.path.join(CREDENCIALES_DIR, nombre_archivo)
    wb.save(ruta)
    return send_file(ruta, as_attachment=True, download_name=nombre_archivo)


@app.route('/api/invitados')
def api_invitados():
    conn = get_db()
    invitados_list = conn.execute('SELECT * FROM invitados ORDER BY id').fetchall()
    conn.close()
    return jsonify([dict(r) for r in invitados_list])


if __name__ == '__main__':
    import sys
    import socket

    hostname = socket.gethostname()
    local_ip = socket.gethostbyname(hostname)

    print(f'\n  Sistema de Registro de Invitados')
    print(f'  ================================')
    print(f'  Local:     http://localhost:5000')
    print(f'  Red:       http://{local_ip}:5000')
    print(f'  Escanear:  http://{local_ip}:5000/escanear')
    print(f'  Dashboard: http://{local_ip}:5000/dashboard')
    print()

    if '--dev' in sys.argv:
        app.run(host='0.0.0.0', port=5000, debug=True)
    else:
        try:
            from waitress import serve
            print('  Servidor: waitress (produccion, multi-usuario)')
            print('  Listo para recibir conexiones...\n')
            serve(app, host='0.0.0.0', port=5000, threads=8)
        except ImportError:
            print('  Servidor: Flask desarrollo (instale waitress para produccion)')
            print('  pip install waitress\n')
            app.run(host='0.0.0.0', port=5000, debug=True)
